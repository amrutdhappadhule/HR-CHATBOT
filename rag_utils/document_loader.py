
from pathlib import Path
import csv
import io


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".pdf", ".docx", ".html", ".htm"
}


def extract_text(filename, content):
    """Extract readable text from a supported file."""

    extension = Path(filename).suffix.lower()

    if extension in {".txt", ".md"}:
        return content.decode("utf-8-sig", errors="replace").strip()

    if extension == ".csv":
        stream = io.StringIO(
            content.decode("utf-8-sig", errors="replace")
        )
        rows = csv.reader(stream)
        return "\n".join(
            " | ".join(cell.strip() for cell in row)
            for row in rows
        ).strip()

    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = []

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()

            if page_text:
                pages.append(
                    f"[Page {page_number}]\n{page_text}"
                )

        text = "\n\n".join(pages).strip()

        if not text:
            raise ValueError(
                "No readable text found. This may be a scanned PDF "
                "that requires OCR."
            )

        return text

    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]

                if any(cells):
                    parts.append(" | ".join(cells))

        return "\n".join(parts).strip()

    if extension in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        html = content.decode("utf-8", errors="replace")
        soup = BeautifulSoup(html, "html.parser")

        for element in soup(["script", "style", "noscript"]):
            element.decompose()

        return soup.get_text(separator="\n", strip=True)

    raise ValueError(f"Unsupported file format: {extension}")
